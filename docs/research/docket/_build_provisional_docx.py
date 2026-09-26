#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Assemble COSMOS provisional pack into one Word document. Not a filing."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"V:\A\Ai\COSMOS")
DOCKET = ROOT / "docs" / "research" / "docket"
RESEARCH = ROOT / "work_orders" / "ccr" / "CREW" / "OUT" / "DOCKET"
DESKTOP = Path(r"C:\Users\Papa\OneDrive\Desktop")
OUT = DESKTOP / "COSMOS_Provisional_Patent_Applications.docx"
OUT2 = DOCKET / "COSMOS_Provisional_Patent_Applications.docx"

FILE_PACKETS = [
    ("P01_MOTIF_LOOP.md", "COSMOS-P01"),
    ("P02_DUAL_LANE.md", "COSMOS-P02"),
    ("P03_POROSITY.md", "COSMOS-P03"),
    ("P04_RUNTIME_BIND.md", "COSMOS-P04"),
    ("P05_ADVERSARIAL_ENGINE.md", "COSMOS-P05"),
    ("P06_LOCAL_FREE_WEIGHTS.md", "COSMOS-P06"),
    ("P07_CORE_LEDGER_FENCE.md", "COSMOS-P07"),
    ("P08_DOM_FIRST.md", "COSMOS-P08"),
    ("P09_SPEND_GATE.md", "COSMOS-P09"),
    ("P10_RESOLVER.md", "COSMOS-P10"),
    ("P11_SEED.md", "COSMOS-P11"),
    ("P13_THINKFAST_DROP.md", "COSMOS-P13"),
]

SUPPORT = [
    DOCKET / "VETTING.md",
    RESEARCH / "R1_OCCUPANCY.md",
    RESEARCH / "R2_LOCAL_FREE.md",
    RESEARCH / "R3_CORE.md",
    RESEARCH / "R4_CRUCIBLE_TM.md",
    RESEARCH / "R5_ORTHOGONAL_POROSITY.md",
    DOCKET / "HOW_IT_WORKS.md",
    DOCKET / "P12_CRUCIBLE.md",
]

INVENTORY = """
## Twenty-five ideas found in the COSMOS universe

These are inventor-named concepts evidenced in the live tree and docket. They are not twenty-five separate $65 filings. Counsel picks claims. This pack files twelve written descriptions (P01–P11 and P13). P12 is HOLD. Remaining ideas are embodiments, skins, or supporting mechanisms — listed so they are not lost.

1. MOTIF nine-stage loop with DEFINE first and ITERATE returning to DEFINE then RESEARCH (P01).
2. Dual-lane BUILD with no shared context; two builders, not builder-plus-checker (P02).
3. Family-axis occupancy / hole-set porosity; different-mistake overlay vs copies (P03).
4. Orthogonal porosity vector and tensor grid T[model_i, model_j, axis]; pair magnitude = disagreement frequency times error magnitude (P03 embodiment / R5).
5. Runtime-binding gate: a value only the live tree can emit; green logs are not done (P04).
6. Adversarial occupancy engine: N isolated seats, different family, one disposer, product skins (P05).
7. Local free-weight plurality vs one S-tier on box-versus-token dollars (P06).
8. Sole signed append-only hash-chained JSONL ledger and fencing tokens (P07).
9. DOM-first unmetered rail when API credit can lapse (P08).
10. Fail-closed spend gate; confirm-to-widen (P09).
11. Sentinel root resolver; existence is not identity (P10).
12. Signed session close (SEED) and OPEN_CONTEXT on improper close (P11).
13. ThinkFast drop: voice or Chatbox to GitHub drop to native OS daemon to named agent to timestamped audit (P13).
14. ChatBot phone Freemium named :free picker (P13 embodiment, not a 13th slot).
15. Remote adversarial occupancy seated from terminal or phone through the same drop (P05 through P13).
16. Crucible legal seats (P12 HOLD — PDJ triad granted in CN and claimed in US20260037351A1).
17. One Chief Coder lease; agents propose, one writer disposes (occupancy / P05 disposer).
18. Two-pen mailbox before any shared tree (BTS two-writer scar).
19. Prompt-cache prefix orthodoxy: identical bytes, volatile last, measure cached_tokens (principle P11 in code; not a 13th filing).
20. Native extra-pane dashboard (cDeck) as product skin on one Core API, not a second Core.
21. Parallel profile windows (Forge and Crucible) on one Core; peer host stays NO_HOST.
22. Gitur BUILD triad (GitHub, GitLab, Cursor) as the branched implementation path, not a fourth writer.
23. CLOCKS collapse: one HOLD-aware Pulse, one claim_next runner, calendar --once.
24. Open Sessions / session-tools suite (list, open, convert, diff, check, anonymize) as first product on the OS.
25. Keep-her-afloat self-modification: the live tree stays up through fenced change; fail-closed refuse rather than silent repair.

Reliable = 2+3+4+5+6. Scalable = 7. OS = 8–12. Loop = 1. Ingress = 13. P12 stays HOLD.
"""

COVER = """
This volume is attorney work product for counsel. It is not a filed patent application, not legal advice, not a novelty opinion, and not a claim set. Inventor legal name is blank for counsel. This preparer does not sign as inventor and does not click USPTO.

Twelve standalone written descriptions are included for US provisional filing under 37 CFR 1.16(d) (micro-entity $65 each if Keith certifies). A provisional cannot claim benefit of a sister. Duplicate the How-It-Works appendix into each filing. Do not add technical matter after a filing date (37 CFR 1.53(c)). No IDS on a provisional.

Public GitHub already started a disclosure clock: bts-mesh 2026-08-16T10:21:14Z; cosmos 2026-08-23T06:42:12Z; cDeck 2026-08-24T18:38:08Z. US inventor grace is about one year from the inventor's disclosure. Many other countries: absolute novelty. Private-now does not un-publish. Do not post whitepaper, X, LinkedIn, or arXiv until counsel files.

Official USPTO TESS and Patent Public Search on DOM rails were not executed in the 2026-09-07 research pass. Prior-art tables below are Google Patents, arXiv, and secondary trademark databases. Combination novelty is UNKNOWN without a claim chart.

Best-mode embodiment: the COSMOS live tree, tree_id KMesh-COSMOS-live, Core on loopback port 8770.
"""


def set_run_font(run, size=12, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_runs(p, text, size=12, bold=False):
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = p.add_run(part[2:-2])
            set_run_font(r, size=size, bold=True)
        else:
            r = p.add_run(part)
            set_run_font(r, size=size, bold=bold)


def add_para(doc, text, *, style=None, size=12, bold=False, space_after=8, align=None):
    p = doc.add_paragraph()
    if style:
        p.style = style
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if align:
        p.alignment = align
    add_runs(p, text, size=size, bold=bold)
    return p


def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size={1: 16, 2: 14, 3: 13}.get(level, 12), bold=True)
    return p


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[j] if j < len(row) else ""
            r = p.add_run(val)
            set_run_font(r, size=10, bold=(i == 0))
    doc.add_paragraph()


def parse_md(doc, text: str):
    lines = text.replace("\r\n", "\n").split("\n")
    i = 0
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        rows = []
        for raw in table_buf:
            cells = [c.strip() for c in raw.strip().strip("|").split("|")]
            if cells and re.match(r"^[-: ]+$", "".join(cells)):
                continue
            rows.append(cells)
        add_table(doc, rows)
        table_buf = []

    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("|") and "|" in line[1:]:
            table_buf.append(line)
            i += 1
            continue
        flush_table()
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            buf = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            add_para(doc, "\n".join(buf)[:4000], size=10, space_after=8)
            i += 1
            continue
        if s == "---":
            i += 1
            continue
        if s.startswith("# "):
            add_heading(doc, s[2:].strip(), 1)
        elif s.startswith("## "):
            add_heading(doc, s[3:].strip(), 2)
        elif s.startswith("### "):
            add_heading(doc, s[4:].strip(), 3)
        elif s.startswith("- "):
            add_para(doc, s[2:].strip(), size=12)
        else:
            add_para(doc, s, size=12)
        i += 1
    flush_table()


def page_break(doc):
    doc.add_page_break()


def main() -> int:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    header = section.header.paragraphs[0]
    hr = header.add_run(
        "COSMOS docket  ·  ATTORNEY WORK PRODUCT — NOT A FILED APPLICATION"
    )
    set_run_font(hr, size=9, italic=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Page ")
    set_run_font(fr, size=9)
    # simple page field omitted; Word will show footer text
    fr2 = footer.add_run("· counsel copy · inventor name blank")
    set_run_font(fr2, size=9)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    add_heading(doc, "COSMOS — Twelve US Provisional Written Descriptions", 1)
    add_para(
        doc,
        "Carry-Over State Mesh Operating System  ·  Docket MOTIF 1  ·  2026-09-09 assembly",
        size=12,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        "Inventor: ________________________________  (counsel to complete)",
        size=12,
    )
    add_para(doc, COVER.strip())

    add_heading(doc, "Part A — Idea inventory and filing map", 1)
    parse_md(doc, INVENTORY.strip())
    add_para(
        doc,
        "The twelve FILE written descriptions in Part C are P01–P11 and P13. "
        "P12 Crucible is reproduced in Part E as HOLD only. ChatBot Freemium and "
        "remote occupancy ride as embodiments on P13 and P05. They are not extra $65 slots.",
    )

    add_heading(doc, "Part B — Prior-art summary (not a novelty opinion)", 1)
    add_para(
        doc,
        "Source: docket VETTING.md and independent research returns R1–R5 "
        "(2026-09-07). Official USPTO TESS and Patent Public Search were not run. "
        "UNKNOWN means not verified, not “does not exist.” Combination novelty of "
        "the twelve packets as a system is UNKNOWN without a claim chart.",
    )
    add_para(
        doc,
        "Closest named hits counsel must read: CN121436020A (eight-stage agent loop); "
        "US20250165890A1 Twilio planner-critic-executor; ChatDev arXiv:2307.07924; "
        "MetaGPT arXiv:2308.00352; US12,405,822 B1 OpenAI shared workspace (anti-shape for P02); "
        "US20260252812A1 IBM contesting LLMs; Hidden Clones arXiv:2603.17111; "
        "Knight & Leveson 1986 N-version; Kuncheva & Whitaker 2003 diversity; "
        "Microsoft US12524210B2 local vs remote coding COGS; Google US20240311405A1 pick-one-of-N; "
        "US11943344B2 signed ledger neighborhood; CONTINUITY arXiv:2609.05269 (session continuity); "
        "CN119168059B and US20260037351A1 PDJ triad (P12 HOLD). "
        "Marks: COSMOS crowded (NVIDIA Cosmos); MOTIF/Forge/Crucible crowded; "
        "cDeck and Gitur looked cleanest on secondary DBs — confirm TESS before $350/class.",
    )
    add_para(
        doc,
        "P01 risk is high if claimed as any staged multi-agent SDLC; lower if DEFINE-verbatim "
        "plus iterate-to-RESEARCH plus dual-lane plus vendor-plural is one method. "
        "P02 must not claim two agents in parallel generally; isolation and two builders "
        "are the thin element. P03 patents are thin; papers are crowded; do not claim Reason "
        "or a fake exponent. P04 is thin in patents if limited to live-emit-only. "
        "P06 must not claim MoA/FrugalGPT slogans. P12 is HOLD.",
    )

    add_heading(doc, "Part C — Twelve FILE written descriptions", 1)
    add_para(
        doc,
        "Each of the following is a standalone specification. Duplicate How-It-Works at filing. "
        "No numbered claim set. Paragraph numbers in the source packets are preserved as text.",
    )
    for name, docket_id in FILE_PACKETS:
        path = DOCKET / name
        page_break(doc)
        add_heading(doc, docket_id + " — " + name.replace(".md", ""), 1)
        parse_md(doc, path.read_text(encoding="utf-8"))

    page_break(doc)
    add_heading(doc, "Part D — Supporting research returns (R1–R5) and vetting", 1)
    add_para(
        doc,
        "Full research occupancy tables. Not exhaustive. Not an IDS. Not a novelty opinion.",
    )
    for path in SUPPORT:
        if not path.is_file():
            add_para(doc, "MISSING: " + path.name)
            continue
        if path.name == "P12_CRUCIBLE.md":
            continue
        if path.name == "HOW_IT_WORKS.md":
            continue
        page_break(doc)
        add_heading(doc, path.name, 1)
        parse_md(doc, path.read_text(encoding="utf-8"))

    page_break(doc)
    add_heading(doc, "Part E — HOLD packet P12 (not the 12th $65 slot)", 1)
    parse_md(doc, (DOCKET / "P12_CRUCIBLE.md").read_text(encoding="utf-8"))

    page_break(doc)
    add_heading(doc, "Part F — How-it-works appendix (attach a copy to each filing)", 1)
    parse_md(doc, (DOCKET / "HOW_IT_WORKS.md").read_text(encoding="utf-8"))

    page_break(doc)
    add_heading(doc, "Part G — What this volume is not", 1)
    add_para(
        doc,
        "This volume does not file anything. It does not complete inventor oath. "
        "It does not draft claims. It does not certify micro-entity. It does not "
        "run TESS. It does not post. Counsel and Keith file. CCr does not click USPTO.",
    )

    DESKTOP.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    doc.save(str(OUT2))
    print("wrote", OUT, "bytes", OUT.stat().st_size)
    print("wrote", OUT2, "bytes", OUT2.stat().st_size)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
