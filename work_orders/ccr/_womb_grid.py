#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""WOMB grid: work orders as rows, 100+ fields as columns."""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

CCR = Path(r"V:\A\Ai\COSMOS\work_orders\ccr")
DROP = Path(r"V:\A\Ai\COSMOS\work_orders\drop")
MASTER = CCR / "WOMB_MASTER.jsonl"
OUT = CCR / "WOMB.xlsx"
CELL_CAP = 32000  # Excel hard limit 32767

COLS = [
    # identity
    "n", "order_id", "kind", "source", "origin_surface", "folder", "state",
    "fifo_ts", "Timestamp", "set_id", "ab_pair_id", "arm", "fail_kind",
    # wish / problem
    "wish_checkbox", "wish_verbatim", "wish_why", "motif_stage",
    "frozen_statement_path", "frozen_live_emit", "already_on_lit",
    "pile_why", "job", "chair",
    # SOP six + translation
    "Agent", "Context source", "Task", "Target & scope", "Output",
    "output_what", "expected_form", "route", "item_do", "item_must_not",
    "item_return", "verify_1", "verify_2", "verify_3",
    "wo_path", "write_path", "windows_filename",
    # legend
    "role", "role_enum", "model", "harness", "wrapper", "skill",
    "tools_allow", "tools_forbid", "first_line", "axis", "partner",
    "hero_pack", "hero_pack_ref", "legend_hash", "pen", "propose_only",
    # enviro / size
    "cwd", "wallet", "sandbox", "pack", "window_tokens", "cache_floor_tokens",
    "cache_ttl_s", "cached_tokens", "prompt_tokens", "MAX_tokens",
    "OPTIMUM_tokens", "cap_tokens", "keep_in_tokens_under",
    "budget_out_usd_per_m", "ctx_is_list", "naked_first",
    # crew 3
    "crew1.Agent", "crew1.tier", "crew1.window", "crew1.MAX", "crew1.OPTIMUM",
    "crew1.tps", "crew1.pack_path", "crew1.mouth_form_ok",
    "crew2.Agent", "crew2.tier", "crew2.window", "crew2.MAX", "crew2.OPTIMUM",
    "crew2.tps", "crew2.pack_path", "crew2.mouth_form_ok",
    "crew3.Agent", "crew3.tier", "crew3.window", "crew3.MAX", "crew3.OPTIMUM",
    "crew3.tps", "crew3.pack_path", "crew3.mouth_form_ok",
    # product / grade (empty until run)
    "DONE", "COMPLETED", "output_path", "output_hash", "empty_output_class",
    "spend", "4C", "Verdict.status", "judge_score_hero", "judge_score_base",
    "judge_winner", "wo_partner", "follow_up_oid", "xfer", "attempt",
    "session_id", "n_obs", "orth_sketch", "mag", "porosity_kind",
    "gitur_pr", "checks.github", "checks.gitlab", "gate_value",
    "lit_write", "baseline_oid", "baseline_agent",
]

ENV_RE = {
    "cwd": re.compile(r"^cwd:\s*(.+)$", re.M),
    "wallet": re.compile(r"^wallet:\s*(.+)$", re.M),
    "window_tokens": re.compile(r"^window_tokens:\s*(\d+)", re.M),
    "cache_floor_tokens": re.compile(r"^cache_floor:\s*(\S+)", re.M),
    "budget_out_usd_per_m": re.compile(r"^budget_out_usd_per_m:\s*(\S+)", re.M),
    "sandbox": re.compile(r"^sandbox:\s*(.+)$", re.M),
    "OPTIMUM_tokens": re.compile(r"^OPTIMUM_tokens:\s*(.+)$", re.M),
}


def clip(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return v
    if isinstance(v, list):
        v = " | ".join(str(x) for x in v)
    else:
        v = str(v)
    if v in ("n/a", "NA", "None", "null"):
        return ""
    if len(v) > CELL_CAP:
        return v[: CELL_CAP - 14] + "…[clipped]"
    return v


def parse_env(text: str) -> dict:
    out = {}
    if not text:
        return out
    for k, rx in ENV_RE.items():
        m = rx.search(text)
        if not m:
            continue
        raw = m.group(1).strip()
        if k in ("window_tokens",):
            out[k] = int(raw)
        elif k in ("budget_out_usd_per_m",):
            try:
                out[k] = float(raw)
            except ValueError:
                out[k] = raw
        elif k == "cache_floor_tokens":
            if raw.isdigit():
                out[k] = int(raw)
            elif raw.lower() in ("n/a", "na"):
                out[k] = ""
            else:
                out[k] = raw
        elif k == "OPTIMUM_tokens":
            out[k] = raw.split()[0] if raw else ""
        else:
            out[k] = raw
    win = out.get("window_tokens")
    if isinstance(win, int) and win > 0:
        out.setdefault("MAX_tokens", int(win * 0.8))
        out.setdefault("cap_tokens", int(win * 0.8))
    return out


def role_enum(role: str) -> str:
    r = (role or "").upper()
    for name in ("WOMBAT", "JUDGE", "CODER", "ORC", "CCR", "DAEMON"):
        if role and (name in r or name in (role or "")):
            return name
    if "proposed patch" in (role or "").lower() or "unified diff" in (role or "").lower():
        return "CODER"
    if "work-order board" in (role or "").lower() or "ITEM or NONE" in (role or ""):
        return "WOMBAT"
    if "KEEP or DROP" in (role or "") or "grade a proposed" in (role or "").lower():
        return "JUDGE"
    return ""


def load_drops() -> dict:
    by = {}
    if not DROP.is_dir():
        return by
    for p in DROP.glob("wo-*.json"):
        try:
            o = json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            continue
        oid = str(o.get("order_id") or p.stem)
        by[oid] = o
    return by


def crew_fill(row: dict, crew: list) -> None:
    for c in crew or []:
        if not isinstance(c, dict):
            continue
        try:
            i = int(c.get("seat") or 0)
        except (TypeError, ValueError):
            i = 0
        if i not in (1, 2, 3):
            continue
        pfx = f"crew{i}."
        row[pfx + "Agent"] = c.get("Agent") or ""
        row[pfx + "tier"] = c.get("tier") or ""
        row[pfx + "window"] = c.get("context_window_tokens") or c.get("window") or ""
        row[pfx + "MAX"] = c.get("MAX_tokens") or ""
        row[pfx + "OPTIMUM"] = c.get("OPTIMUM_tokens") or ""
        row[pfx + "tps"] = c.get("tps") or ""
        row[pfx + "pack_path"] = c.get("pack_path") or c.get("pack") or ""
        mf = c.get("mouth_form_ok")
        row[pfx + "mouth_form_ok"] = mf if mf is not None else ""


def from_master(o: dict) -> dict:
    row = {k: "" for k in COLS}
    row["n"] = o.get("n") or ""
    row["order_id"] = o.get("order_id") or ""
    row["kind"] = o.get("kind") or ""
    row["source"] = o.get("source") or ""
    row["origin_surface"] = o.get("source") or o.get("kind") or ""
    row["folder"] = o.get("folder") or ""
    row["state"] = o.get("state") or ""
    row["fifo_ts"] = o.get("fifo_ts") or ""
    row["Timestamp"] = o.get("Timestamp") or o.get("fifo_ts") or ""
    row["set_id"] = o.get("set_id") or o.get("order_id") or ""
    row["ab_pair_id"] = o.get("ab_pair_id") or ""
    row["arm"] = o.get("arm") or ""
    row["fail_kind"] = o.get("fail_kind") or ""
    row["Agent"] = o.get("agent_field") or o.get("Agent") or ""
    row["Context source"] = o.get("context_source") or o.get("Context source") or ""
    row["Task"] = o.get("task") or o.get("Task") or ""
    row["Target & scope"] = o.get("target_scope") or o.get("Target & scope") or ""
    row["Output"] = o.get("output_spec") or o.get("Output") or ""
    row["output_what"] = o.get("output_what") or ""
    row["wo_path"] = o.get("wo_path") or ""
    row["write_path"] = o.get("write_path") or ""
    wp = str(row["wo_path"])
    row["windows_filename"] = Path(wp).name if wp else ""
    row["role"] = o.get("role") or ""
    row["role_enum"] = role_enum(str(o.get("role") or ""))
    row["model"] = o.get("model") or ""
    row["harness"] = o.get("harness") or ""
    row["wrapper"] = o.get("wrapper") or ""
    row["skill"] = o.get("skill") or ""
    row["tools_allow"] = o.get("tools_allow") or ""
    row["tools_forbid"] = o.get("tools_forbid") or ""
    row["first_line"] = o.get("first_line") or ""
    row["axis"] = o.get("axis") or ""
    row["partner"] = o.get("partner") or ""
    row["hero_pack"] = o.get("hero_pack") or ""
    row["hero_pack_ref"] = o.get("hero_pack_ref") or ""
    row["pile_why"] = o.get("pile_why") or ""
    row["baseline_oid"] = o.get("baseline_oid") or ""
    row["baseline_agent"] = o.get("baseline_agent") or ""
    row["judge_score_hero"] = o.get("judge_score_hero") if o.get("judge_score_hero") not in (None, "n/a") else ""
    row["judge_score_base"] = o.get("judge_score_base") if o.get("judge_score_base") not in (None, "n/a") else ""
    row["judge_winner"] = o.get("judge_winner") or ""
    row["chair"] = "WOMBAT"
    row["pen"] = "none"
    row["propose_only"] = True
    row["lit_write"] = False
    row["ctx_is_list"] = True
    row["naked_first"] = False
    row["DONE"] = False
    row["COMPLETED"] = False
    row["attempt"] = 0
    row["cache_ttl_s"] = 1800
    row["keep_in_tokens_under"] = 200000
    row["pack"] = "house"
    row.update(parse_env(str(o.get("environment") or "")))
    if not row.get("OPTIMUM_tokens"):
        row["OPTIMUM_tokens"] = "float"
    kind = str(o.get("kind") or "")
    task = str(row["Task"] or "")
    if kind == "wish":
        row["wish_checkbox"] = "- [ ]"
        w = task
        if w.startswith("WISH:"):
            w = w[5:].strip()
        row["wish_verbatim"] = w
        row["job"] = (w.split("—")[0].split(".")[0].strip()[:80] if w else "")
        row["motif_stage"] = "1 PROBLEM STATEMENT"
        row["expected_form"] = "NONE or diff --git"
        row["route"] = "coding"
        row["verify_1"] = "First line NONE or diff --git"
        row["verify_2"] = "Output file exists at write_path"
        row["verify_3"] = "No LiT write; no grok.exe worker"
        row["item_must_not"] = "LiT write; grok.exe; USPTO; 112x MOTIF"
        row["item_return"] = "unified diff first"
    if "AUTO-RESESSION" in task.upper() or "AUTO-RESESSION" in str(row.get("wish_verbatim") or "").upper():
        row["wish_why"] = (
            "Layer B orch dies when the window fills; Layer A clocks keep running. "
            "Need a satellite that spawns a FRESH orch from SEED/BU. --continue fails the wish."
        )
        row["frozen_statement_path"] = r"V:\A\Ai\COSMOS\docs\research\AUTO_RESESSION.md"
        row["frozen_live_emit"] = "RESESSION.json + resession_heartbeat.json + clock 18"
        row["already_on_lit"] = "cosmos_resession.py CLOCK_ID=18; --install-task still operator"
        row["item_do"] = "Satellite spawn of fresh orch; not resume of dying window"
        row["job"] = "AUTO-RESESSION"
    # produce cells stay empty / false / 0 until run — not UNMEASURED text on every row
    return row


def overlay_drop(row: dict, d: dict) -> None:
    if not d:
        return
    for src, dst in (
        ("Agent", "Agent"),
        ("Task", "Task"),
        ("Target & scope", "Target & scope"),
        ("Output", "Output"),
        ("output_what", "output_what"),
        ("Timestamp", "Timestamp"),
        ("MAX_tokens", "MAX_tokens"),
        ("OPTIMUM_tokens", "OPTIMUM_tokens"),
        ("Context source", "Context source"),
        ("order_id", "order_id"),
    ):
        if d.get(src) not in (None, ""):
            row[dst] = d[src]
    crew_fill(row, d.get("crew") or [])
    if row.get("crew1.MAX") and not row.get("MAX_tokens"):
        row["MAX_tokens"] = row["crew1.MAX"]
    if row.get("crew1.window") and not row.get("window_tokens"):
        row["window_tokens"] = row["crew1.window"]


def main() -> int:
    assert len(COLS) == len(set(COLS)), "duplicate columns"
    drops = load_drops()
    rows = []
    with MASTER.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            o = json.loads(line)
            row = from_master(o)
            oid = str(row.get("order_id") or "")
            overlay_drop(row, drops.get(oid))
            rows.append(row)

    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = "WOMB"
    ws.append(COLS)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "C2"
    for row in rows:
        ws.append([clip(row.get(k, "")) for k in COLS])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(rows)+1}"
    # modest widths
    for i, name in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(10, len(name) + 2))
    meta = wb.create_sheet("Shape")
    meta.append(["rows", "columns", "cells", "meaning"])
    meta.append([len(rows), len(COLS), len(rows) * len(COLS), "jobs x fields"])
    meta.append(["thin master keys", 41, "", "WOMB_MASTER.jsonl"])
    meta.append(["this grid", len(COLS), "", "100+ columns; produce cells empty until run"])
    meta.append(["catalog", 266, "", "WOMB_ULTIMATE.xlsx is the field dictionary, not the board"])
    wb.save(OUT)
    print(json.dumps({
        "ok": True,
        "path": str(OUT),
        "rows": len(rows),
        "columns": len(COLS),
        "cells": len(rows) * len(COLS),
        "drops_overlaid": sum(1 for r in rows if r["order_id"] in drops),
    }))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
